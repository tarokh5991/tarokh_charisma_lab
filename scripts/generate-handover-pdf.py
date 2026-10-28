#!/usr/bin/env python3
from pathlib import Path

import arabic_reshaper
from bidi.algorithm import get_display
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "project-handover.md"
OUTPUT = ROOT / "outputs" / "tarokh-project-handover-1405-07-28.pdf"
SOURCE_DOCUMENTS = (
    ("استراتژی مخاطب پیج کاریزما", ROOT / "uploads" / "استراتژی-مخاطب-پیج-کاریزما_d5447fd1.docx"),
    ("برندینگ اولیه پیج کاریزما", ROOT / "uploads" / "برندینگ-اولیه-پیج-کاریزما_164d7adc.docx"),
    ("پریزم هویت برند پیج کاریزما", ROOT / "uploads" / "پریزم-هویت-برند-پیج-کاریزما_58959d6b.docx"),
    ("بریف محصول TAROKH", ROOT / "uploads" / "بریف-محصول-TAROKH-سند-اصلی_6cb0f111.docx"),
)

PAGE_WIDTH, PAGE_HEIGHT = 1240, 1754
MARGIN_X, MARGIN_TOP, MARGIN_BOTTOM = 105, 120, 100
CONTENT_WIDTH = PAGE_WIDTH - (2 * MARGIN_X)
BACKGROUND = "#090909"
TEXT = "#f5f1e8"
MUTED = "#a7a399"
ACCENT = "#d6ad5c"
LINE = "#30302c"
FONT = "/usr/share/fonts/truetype/vazirmatn/Vazirmatn-Regular.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/vazirmatn/Vazirmatn-Bold.ttf"


def rtl(value: str) -> str:
    return get_display(arabic_reshaper.reshape(value))


def font(size: int, bold: bool = False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT, size)


def line_height(current_font) -> int:
    box = current_font.getbbox("آبپژ")
    return box[3] - box[1] + 13


def wrap(text: str, current_font, width: int) -> list[str]:
    words = text.split()
    if not words:
        return [""]
    lines: list[str] = []
    current: list[str] = []
    probe = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    for word in words:
        candidate = " ".join(current + [word])
        if current and probe.textlength(rtl(candidate), font=current_font) > width:
            lines.append(" ".join(current))
            current = [word]
        else:
            current.append(word)
    if current:
        lines.append(" ".join(current))
    return lines


class PdfPages:
    def __init__(self):
        self.pages: list[Image.Image] = []
        self.page_number = 0
        self.new_page()

    def new_page(self):
        self.page_number += 1
        self.image = Image.new("RGB", (PAGE_WIDTH, PAGE_HEIGHT), BACKGROUND)
        self.draw = ImageDraw.Draw(self.image)
        self.y = MARGIN_TOP
        self.draw.line((MARGIN_X, 76, PAGE_WIDTH - MARGIN_X, 76), fill=LINE, width=2)
        self.draw.text((PAGE_WIDTH - MARGIN_X, 44), rtl("پروندهٔ تحویل پروژه TAROKH Charisma Lab"), font=font(18), fill=ACCENT, anchor="ra")

    def finish_page(self):
        self.draw.line((MARGIN_X, PAGE_HEIGHT - 64, PAGE_WIDTH - MARGIN_X, PAGE_HEIGHT - 64), fill=LINE, width=2)
        self.draw.text((MARGIN_X, PAGE_HEIGHT - 45), rtl("محرمانه — سند ادامهٔ پروژه"), font=font(16), fill=MUTED)
        self.draw.text((PAGE_WIDTH - MARGIN_X, PAGE_HEIGHT - 45), str(self.page_number), font=font(16), fill=MUTED, anchor="ra")
        self.pages.append(self.image)

    def ensure(self, height: int):
        if self.y + height > PAGE_HEIGHT - MARGIN_BOTTOM:
            self.finish_page()
            self.new_page()

    def add_block(self, text: str, current_font, fill: str = TEXT, indent: int = 0, gap_after: int = 15):
        lines = wrap(text, current_font, CONTENT_WIDTH - indent)
        height = (line_height(current_font) * len(lines)) + gap_after
        self.ensure(height)
        for line in lines:
            self.draw.text((PAGE_WIDTH - MARGIN_X, self.y), rtl(line), font=current_font, fill=fill, anchor="ra")
            self.y += line_height(current_font)
        self.y += gap_after

    def add_heading(self, text: str, level: int):
        size = {1: 36, 2: 28, 3: 23}[level]
        heading_font = font(size, True)
        self.ensure(line_height(heading_font) + 30)
        if level == 1:
            self.y += 12
        self.draw.text((PAGE_WIDTH - MARGIN_X, self.y), rtl(text), font=heading_font, fill=ACCENT, anchor="ra")
        self.y += line_height(heading_font) + (18 if level == 1 else 12)


def source_blocks(document: Document):
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            text = Paragraph(child, document).text.strip()
            if text:
                yield "paragraph", text
        elif isinstance(child, CT_Tbl):
            table = Table(child, document)
            yield "table", [
                [cell.text.strip().replace("\n", " / ") for cell in row.cells]
                for row in table.rows
            ]


def append_source_documents(pages: PdfPages):
    pages.add_heading("۱۳. پیوست‌ها — متن کامل اسناد مبنا", 1)
    pages.add_block(
        "چهار پیوست زیر بدون حذف یا تلخیص به این پرونده افزوده شده‌اند. جدول‌ها به‌صورت سطرهای خوانا بازچینش شده‌اند و متن تمام سلول‌ها حفظ شده است.",
        font(21),
    )
    for index, (title, path) in enumerate(SOURCE_DOCUMENTS, start=1):
        pages.add_heading(f"پیوست {index}: {title}", 2)
        pages.add_block("متن کامل سند مبنا", font(18, True), fill=ACCENT)
        for kind, content in source_blocks(Document(path)):
            if kind == "paragraph":
                pages.add_block(content, font(21), gap_after=11)
            else:
                pages.add_heading("جدول", 3)
                for row in content:
                    pages.add_block(" | ".join(row), font(19), fill=TEXT, gap_after=8)


def render():
    pages = PdfPages()
    for raw in SOURCE.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            pages.y += 6
            continue
        if line.startswith("# "):
            pages.add_heading(line[2:], 1)
        elif line.startswith("## "):
            pages.add_heading(line[3:], 2)
        elif line.startswith("### "):
            pages.add_heading(line[4:], 3)
        elif line.startswith("- "):
            pages.add_block("• " + line[2:], font(21), indent=16, gap_after=7)
        elif len(line) > 2 and line[0].isdigit() and ". " in line[:5]:
            pages.add_block(line, font(21), indent=16, gap_after=7)
        else:
            pages.add_block(line.replace("**", ""), font(21), gap_after=15)
    append_source_documents(pages)
    pages.finish_page()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    pages.pages[0].save(
        OUTPUT,
        "PDF",
        resolution=150.0,
        save_all=True,
        append_images=pages.pages[1:],
        title="TAROKH Charisma Lab Project Handover",
        author="TAROKH Charisma Lab",
    )
    print(OUTPUT)
    print(f"pages={len(pages.pages)}")


if __name__ == "__main__":
    render()
